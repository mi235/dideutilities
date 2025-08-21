import os
from tkinter import Tk, filedialog
from openpyxl import load_workbook

# Διάλεξε Excel
excel_file = filedialog.askopenfilename(
    title="Επίλεξε το Excel",
    filetypes=[("Excel Files", "*.xlsx")]
)

# Διάλεξε τον βασικό φάκελο
root = Tk()
root.withdraw()  # κρύβει το κεντρικό παράθυρο
base_dir = filedialog.askdirectory(title="Επίλεξε τον βασικό φάκελο")
root.destroy()

# Φόρτωσε το Excel
wb = load_workbook(excel_file)
sheet = wb.active

# Δημιουργία υποφακέλων
for row in sheet.iter_rows(min_row=2, values_only=True):  # από τη 2η γραμμή
    folder_name = str(row[0]).strip()
    if folder_name:
        folder_path = os.path.join(base_dir, folder_name)
        os.makedirs(folder_path, exist_ok=True)
        print(f"Δημιουργήθηκε: {folder_path}")
