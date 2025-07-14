import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def eisagogi_grammon_klados_me_epikefalida(df):
    # Αντικατάσταση κενών στην ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ
    df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"] = df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"].replace("", pd.NA)
    df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"] = df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"].fillna(df["ΑΠΟ"])

    # Αφαίρεση της στήλης "ΑΠΟ"
    if "ΑΠΟ" in df.columns:
        df = df.drop(columns=["ΑΠΟ"])

    # Ταξινόμηση
    df = df.sort_values(["KLADOS", "ΕΙΔΙΚΗ ΚΑΤΗΓΟΡΙΑ", "MORIA"], ascending=[True, True, False])
    columns = df.columns.tolist()
    rows = []

    for klados, group in df.groupby("KLADOS"):
        empty_row = [""] * len(columns)
        empty_row[0] = f"ΚΛΑΔΟΣ = {klados}"
        rows.append(empty_row)
        rows.append(columns)
        rows.extend(group.values.tolist())

    return pd.DataFrame(rows, columns=columns)

def export_with_formatting(df_final, output_path):
    df_final.to_excel(output_path, index=False, sheet_name="ΑΠΟΤΕΛΕΣΜΑ")

    wb = load_workbook(output_path)
    ws = wb["ΑΠΟΤΕΛΕΣΜΑ"]

    # Προσδιορισμός αριθμών στηλών
    col_map = {cell.value: idx + 1 for idx, cell in enumerate(ws[1]) if cell.value}

    for row in ws.iter_rows(min_row=1):
        first_cell = row[0]
        if isinstance(first_cell.value, str) and first_cell.value.startswith("ΚΛΑΔΟΣ ="):
            first_cell.font = Font(size=24, bold=True)

        # Στήλη "ΥΠΑΛΛΗΛΟΣ"
        if "ΥΠΑΛΛΗΛΟΣ" in col_map:
            cell = row[col_map["ΥΠΑΛΛΗΛΟΣ"] - 1]
            if cell.value:
                cell.font = Font(size=12)

        # Αναδίπλωση κειμένου
        for col_title in ["ΠΡΟΤΙΜΗΣΕΙΣ ΣΧΟΛΕΙΩΝ", "ΣΧΟΛΙΑ- ΛΟΓΟΙ ΤΟΠΟΘΕΤΗΣΗΣ"]:
            if col_title in col_map:
                cell = row[col_map[col_title] - 1]
                cell.alignment = Alignment(wrap_text=True)

    # Αυτόματο πλάτος στηλών
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    wb.save(output_path)

# --- Χρήση ---
df = pd.read_excel("MORIA.xlsx")
df_final = eisagogi_grammon_klados_me_epikefalida(df)
export_with_formatting(df_final, "MORIA_morfopoiimeno.xlsx")
