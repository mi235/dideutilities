import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

def epilogi_eisodou():
    root = tk.Tk()
    root.withdraw()
    filepath = filedialog.askopenfilename(
        title="Επιλογή αρχείου Excel",
        filetypes=[("Excel αρχεία", "*.xlsx *.xls")]
    )
    return filepath

def epilogi_exodou():
    root = tk.Tk()
    root.withdraw()
    filepath = filedialog.asksaveasfilename(
        title="Αποθήκευση νέου αρχείου Excel",
        defaultextension=".xlsx",
        filetypes=[("Excel αρχεία", "*.xlsx")]
    )
    return filepath

def eisagogi_grammon_klados_me_epikefalida(df):
    # Αντικατάσταση κενών στην ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ με τιμή από "ΑΠΟ"
    df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"] = df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"].replace("", pd.NA)
    df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"] = df["ΠΕΡΙΟΧΗ ΜΕΤΑΘΕΣΗΣ"].fillna(df["ΑΠΟ"])

    # Διαγραφή της στήλης "ΑΠΟ"
    if "ΑΠΟ" in df.columns:
        df = df.drop(columns=["ΑΠΟ"])

    # Ταξινόμηση
    df = df.sort_values(["KLADOS", "ΕΙΔΙΚΗ ΚΑΤΗΓΟΡΙΑ", "MORIA"], ascending=[True, True, False])
    columns = df.columns.tolist()
    rows = []

    for klados, group in df.groupby("KLADOS"):
        # Γραμμή με ΚΛΑΔΟΣ
        klados_row = [""] * len(columns)
        klados_row[0] = f"ΚΛΑΔΟΣ = {klados}"
        rows.append(klados_row)

        # Επικεφαλίδα στηλών
        rows.append(columns)

        # Προσθήκη δεδομένων
        rows.extend(group.values.tolist())

    return pd.DataFrame(rows, columns=columns)

def export_with_formatting(df_final, output_path):
    df_final.to_excel(output_path, index=False, sheet_name="ΑΠΟΤΕΛΕΣΜΑ")

    wb = load_workbook(output_path)
    ws = wb["ΑΠΟΤΕΛΕΣΜΑ"]

    # Εύρεση αριθμού στηλών βάσει επικεφαλίδας
    col_map = {cell.value: idx + 1 for idx, cell in enumerate(ws[1]) if cell.value}

    for row in ws.iter_rows(min_row=1):
        first_cell = row[0]
        if isinstance(first_cell.value, str) and first_cell.value.startswith("ΚΛΑΔΟΣ ="):
            first_cell.font = Font(size=24, bold=True)

        # Γραμματοσειρά 12 στη στήλη ΥΠΑΛΛΗΛΟΣ
        if "ΥΠΑΛΛΗΛΟΣ" in col_map:
            cell = row[col_map["ΥΠΑΛΛΗΛΟΣ"] - 1]
            if cell.value:
                cell.font = Font(size=12)

        # Αναδίπλωση κειμένου σε συγκεκριμένες στήλες
        for col_title in ["ΠΡΟΤΙΜΗΣΕΙΣ ΣΧΟΛΕΙΩΝ", "ΣΧΟΛΙΑ- ΛΟΓΟΙ ΤΟΠΟΘΕΤΗΣΗΣ"]:
            if col_title in col_map:
                cell = row[col_map[col_title] - 1]
                cell.alignment = Alignment(wrap_text=True)

    # Αυτόματο πλάτος στηλών
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 2

    wb.save(output_path)

def main():
    # Επιλογή αρχείου εισόδου
    excel_file = epilogi_eisodou()
    if not excel_file:
        messagebox.showwarning("Προσοχή", "Δεν επιλέχθηκε αρχείο εισόδου.")
        return

    # Ανάγνωση αρχείου
    try:
        df = pd.read_excel(excel_file)
    except Exception as e:
        messagebox.showerror("Σφάλμα", f"Σφάλμα κατά την ανάγνωση:\n{e}")
        return

    # Επεξεργασία
    new_df = eisagogi_grammon_klados_me_epikefalida(df)

    # Επιλογή αρχείου εξόδου
    output_file = epilogi_exodou()
    if not output_file:
        messagebox.showinfo("Ακύρωση", "Δεν επιλέχθηκε αρχείο εξόδου.")
        return

    # Αποθήκευση με μορφοποίηση
    try:
        export_with_formatting(new_df, output_file)
        messagebox.showinfo("Ολοκληρώθηκε", f"Το αρχείο αποθηκεύτηκε:\n{output_file}")
    except Exception as e:
        messagebox.showerror("Σφάλμα", f"Σφάλμα κατά την αποθήκευση:\n{e}")

if __name__ == "__main__":
    main()
